import os
import shutil
from pathlib import Path
import pptx
from pptx import Presentation

from app.config import DATA_DIR

MSO_TRUE = -1
MSO_FALSE = 0
MSO_MEDIA = 16


class PowerPointService:
    def __init__(self):
        self.thumbs_dir = DATA_DIR / "thumbnails"
        self.thumbs_dir.mkdir(parents=True, exist_ok=True)

    def extract_slides_info(self, pptx_path: str | Path, generate_thumbnails: bool = True) -> list[dict]:
        pptx_path = str(Path(pptx_path).resolve())
        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"Presentation not found: {pptx_path}")

        # 1. Parse text & notes via python-pptx
        prs = Presentation(pptx_path)
        slides_data = []

        for idx, slide in enumerate(prs.slides):
            slide_number = idx + 1
            title = ""
            body_text = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        if not title and (shape == slide.shapes[0] or getattr(shape, "is_title", False)):
                            title = text
                        else:
                            body_text.append(text)

            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()

            slides_data.append({
                "slide_number": slide_number,
                "title": title or f"Slide {slide_number}",
                "body_text": "\n".join(body_text),
                "notes": notes_text,
                "thumbnail_path": "",
            })

        # 2. Export slide thumbnails via PowerPoint COM if available
        if generate_thumbnails:
            try:
                import win32com.client
                ppt = win32com.client.Dispatch("PowerPoint.Application")
                pres = ppt.Presentations.Open(pptx_path, ReadOnly=True, Untitled=False, WithWindow=False)

                pptx_stem = Path(pptx_path).stem
                for idx, slide_item in enumerate(slides_data):
                    slide_num = slide_item["slide_number"]
                    thumb_file = self.thumbs_dir / f"{pptx_stem}_slide_{slide_num}.png"
                    try:
                        com_slide = pres.Slides(slide_num)
                        com_slide.Export(str(thumb_file.resolve()), "PNG", 320, 180)
                        slide_item["thumbnail_path"] = str(thumb_file.resolve())
                    except Exception:
                        pass
                pres.Close()
            except Exception:
                # COM thumbnail export optional fallback
                pass

        return slides_data

    def embed_audio_to_slide(
        self,
        pres_com,
        slide_number: int,
        audio_file_path: str | Path,
        autoplay: bool = True,
        hide_icon: bool = True,
        x_pos: int = 15,
        y_pos: int = 15,
    ):
        audio_file = str(Path(audio_file_path).resolve())
        slide = pres_com.Slides(slide_number)

        # Remove previous voice studio audio shapes
        to_delete = []
        for i in range(1, slide.Shapes.Count + 1):
            try:
                shp = slide.Shapes(i)
                if shp.Name.startswith("VoiceStudioAudio"):
                    to_delete.append(shp)
            except Exception:
                pass

        for shp in to_delete:
            try:
                shp.Delete()
            except Exception:
                pass

        # Embed audio
        # AddMediaObject2(FileName, LinkToFile, SaveWithDocument, Left, Top, Width, Height)
        shape = slide.Shapes.AddMediaObject2(audio_file, MSO_FALSE, MSO_TRUE, x_pos, y_pos, 40, 40)
        shape.Name = f"VoiceStudioAudio_{slide_number}"

        if autoplay:
            shape.AnimationSettings.PlaySettings.PlayOnEntry = MSO_TRUE
            shape.AnimationSettings.Animate = MSO_TRUE

        if hide_icon:
            shape.AnimationSettings.PlaySettings.HideWhileNotPlaying = MSO_TRUE

        return shape

    def embed_video_to_slide(
        self,
        pres_com,
        slide_number: int,
        video_file_path: str | Path,
        left: int = 80,
        top: int = 80,
        width: int = 800,
        height: int = 450,
        autoplay: bool = False,
    ):
        video_file = str(Path(video_file_path).resolve())
        slide = pres_com.Slides(slide_number)

        # Remove previous voice studio video shapes on this slide
        to_delete = []
        for i in range(1, slide.Shapes.Count + 1):
            try:
                shp = slide.Shapes(i)
                if shp.Name.startswith("VoiceStudioVideo"):
                    to_delete.append(shp)
            except Exception:
                pass

        for shp in to_delete:
            try:
                shp.Delete()
            except Exception:
                pass

        shape = slide.Shapes.AddMediaObject2(video_file, MSO_FALSE, MSO_TRUE, left, top, width, height)
        shape.Name = f"VoiceStudioVideo_{slide_number}"

        if autoplay:
            shape.AnimationSettings.PlaySettings.PlayOnEntry = MSO_TRUE
            shape.AnimationSettings.Animate = MSO_TRUE

        return shape

    def batch_export_presentation(
        self,
        source_pptx: str | Path,
        output_pptx: str | Path,
        slide_audio_map: dict[int, str],
        slide_video_map: dict[int, str] | None = None,
        progress_callback=None,
    ) -> str:
        source_pptx = str(Path(source_pptx).resolve())
        output_pptx = str(Path(output_pptx).resolve())

        # Copy original first to avoid altering source
        shutil.copy2(source_pptx, output_pptx)

        import win32com.client
        import time
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        # WithWindow=True is required by PowerPoint media engine to commit AddMediaObject2 streams
        pres = ppt.Presentations.Open(output_pptx, ReadOnly=False, Untitled=False, WithWindow=True)

        total_slides = pres.Slides.Count
        try:
            for slide_num in range(1, total_slides + 1):
                if progress_callback:
                    progress_callback(slide_num, total_slides, f"Processing slide {slide_num}/{total_slides}...")

                # Embed audio if present for this slide
                if slide_num in slide_audio_map and os.path.exists(slide_audio_map[slide_num]):
                    self.embed_audio_to_slide(pres, slide_num, slide_audio_map[slide_num], autoplay=True)

                # Embed video if present for this slide
                if slide_video_map and slide_num in slide_video_map and os.path.exists(slide_video_map[slide_num]):
                    self.embed_video_to_slide(pres, slide_num, slide_video_map[slide_num], autoplay=False)

            pres.Save()
            time.sleep(0.3)
        finally:
            try:
                pres.Close()
            except Exception:
                pass

        return output_pptx
